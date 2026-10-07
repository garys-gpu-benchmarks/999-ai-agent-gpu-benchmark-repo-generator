#!/usr/bin/env python3
"""Regenerate Workload_Definitions command columns from the Parameters_SmokeBaselineExtend sheet."""

from __future__ import annotations

import argparse
from pathlib import Path
import sys

from openpyxl import load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

ROOT = Path(__file__).resolve().parents[1]
if str(Path(__file__).resolve().parent) not in sys.path:
    sys.path.insert(0, str(Path(__file__).resolve().parent))

from matrix_profile_values import (  # noqa: E402
    COMMAND_COLUMNS,
    DEFINITIONS_SHEET_ALIASES,
    PROFILE_HEADERS,
    PROFILE_SHEET,
    PROFILE_SHEET_ALIASES,
    PROFILES,
    build_command,
    definitions_rows,
    header_map,
    load_profile_values,
    resolve_sheet,
)


DEFAULT_MATRIX = ROOT / "BenchmarkSpecDefinitions.xlsx"


def parse_args(argv: list[str]) -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Write Workload_Definitions run_benchmark command columns from Parameters_SmokeBaselineExtend."
    )
    parser.add_argument("--matrix", type=Path, default=DEFAULT_MATRIX)
    parser.add_argument(
        "--check",
        action="store_true",
        help="Verify command columns match Parameters_SmokeBaselineExtend without writing.",
    )
    return parser.parse_args(argv)


def ensure_profile_sheet(workbook) -> Worksheet:
    for name in PROFILE_SHEET_ALIASES:
        if name in workbook.sheetnames:
            return workbook[name]
    sheet = workbook.create_sheet(PROFILE_SHEET)
    bold = Font(bold=True)
    for name, column in PROFILE_HEADERS.items():
        cell = sheet.cell(1, column, name)
        cell.font = bold
    sheet.freeze_panes = "A2"
    widths = {1: 16, 2: 36, 3: 24, 4: 48, 5: 48, 6: 48, 7: 56}
    for column, width in widths.items():
        sheet.column_dimensions[get_column_letter(column)].width = width
    return sheet


def write_profile_rows(sheet: Worksheet, rows: list[tuple]) -> None:
    sheet.delete_rows(2, sheet.max_row)
    bold = Font(bold=True)
    for name, column in PROFILE_HEADERS.items():
        cell = sheet.cell(1, column, name)
        cell.font = bold
    for index, row in enumerate(rows, start=2):
        for column, value in enumerate(row, start=1):
            sheet.cell(index, column, value)


def command_column_indexes(definitions_sheet) -> dict[str, int]:
    headers = header_map(definitions_sheet)
    missing = [name for name in COMMAND_COLUMNS.values() if name not in headers]
    if missing:
        raise SystemExit(f"[FAIL] {definitions_sheet.title} missing command columns: {', '.join(missing)}")
    return {profile: headers[COMMAND_COLUMNS[profile]] for profile in PROFILES}


def expected_commands(definitions_sheet, values_by_workload) -> dict[int, dict[str, str]]:
    expected: dict[int, dict[str, str]] = {}
    for workload, info in definitions_rows(definitions_sheet).items():
        values = values_by_workload.get(workload)
        if not values:
            raise SystemExit(f"[FAIL] {PROFILE_SHEET} has no rows for workload {workload}.")
        expected[workload] = {
            profile: build_command(profile, info["parameters"], values) for profile in PROFILES
        }
    return expected


def main(argv: list[str]) -> int:
    args = parse_args(argv)
    workbook = load_workbook(args.matrix)
    definitions_sheet = resolve_sheet(workbook, DEFINITIONS_SHEET_ALIASES, "definitions")
    values_sheet = ensure_profile_sheet(workbook)
    values = load_profile_values(values_sheet)
    if not values:
        raise SystemExit(f"[FAIL] {values_sheet.title} is empty. Author parameter values there first.")
    expected = expected_commands(definitions_sheet, values)
    columns = command_column_indexes(definitions_sheet)
    mismatches: list[str] = []
    for info_workload, info in definitions_rows(definitions_sheet).items():
        row = info["row"]
        for profile, column in columns.items():
            current = definitions_sheet.cell(row, column).value
            wanted = expected[info_workload][profile]
            if str(current or "") != wanted:
                mismatches.append(f"{info_workload} {profile}")
            if not args.check:
                definitions_sheet.cell(row, column).value = wanted
    if args.check:
        if mismatches:
            raise SystemExit(
                f"[FAIL] Command columns do not match {values_sheet.title}: " + ", ".join(mismatches)
            )
        print(f"[PASS] Command columns match {values_sheet.title} for {len(expected)} workloads.")
        return 0
    workbook.save(args.matrix)
    print(f"Wrote command columns for {len(expected)} workloads in {args.matrix}")
    if mismatches:
        print(f"Updated {len(mismatches)} stale command cells.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
