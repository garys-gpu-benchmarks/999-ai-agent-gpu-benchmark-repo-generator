#!/usr/bin/env python3
"""
File: generate_workload_parameters.py
Version: 1.0.0
Author: Cursor
Date: 2026-08-05
Description: Generate the shared workload-parameter catalog from the benchmark matrix workbook.
Execution: python scripts/generate_workload_parameters.py [--check]
Options: --check validates active workbook parameters without writing the catalog.
Requirements: Python 3.10+, openpyxl, PyYAML.
Environment: Local repository workspace.
Dependencies: BenchmarkSpecDefinitions.xlsx (Parameter_* columns by header name).
Variables: WORKBOOK, OUTPUT.
Repository: ai-agent-gpu-benchmark-repo-generator.
License: Project repository license.
"""

from __future__ import annotations

from collections import defaultdict
from pathlib import Path
import re

import yaml
from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "BenchmarkSpecDefinitions.xlsx"
OUTPUT = ROOT / "config" / "workload_parameters.yaml"


def _active_sheet(workbook):
    for name in ("Workload_Definitions", "Sheet4", "Sheet1"):
        if name in workbook.sheetnames:
            return workbook[name]
    return workbook[workbook.sheetnames[0]]


def _header_map(worksheet) -> dict[str, int]:
    return {
        str(worksheet.cell(1, column).value).strip(): column
        for column in range(1, worksheet.max_column + 1)
        if worksheet.cell(1, column).value not in (None, "")
    }


def _parameter_columns(worksheet) -> list[int]:
    headers = _header_map(worksheet)
    columns = [
        column
        for name, column in headers.items()
        if re.fullmatch(r"Parameter_\d+", name)
    ]
    if not columns:
        raise SystemExit("[FAIL] No Parameter_* header columns found in the workbook.")
    return columns


def _description_column(worksheet) -> int:
    headers = _header_map(worksheet)
    if "Execution Description With Parameters" in headers:
        return headers["Execution Description With Parameters"]
    raise SystemExit("[FAIL] Missing 'Execution Description With Parameters' header.")


CATALOG_METADATA = {
    "num_iterations": ("integer", "iterations", "Number of timed repetitions.", "positive integer", "termination control; may interact with duration"),
    "dtype": ("string", "format token", "Tensor, element, or storage data type.", "workload-specific supported dtype token", "keep distinct from precision when storage and compute formats differ"),
    "warmup_iters": ("integer", "iterations", "Untimed warm-up repetitions before measurement.", "integer >= 0", "excluded from measured iteration counts"),
    "batch_size": ("integer", "items per batch", "Number of items processed together.", "integer >= 1", "constrained by available memory and sometimes num_gpus"),
    "model_name": ("string", "model identifier", "Model or weights identifier to load.", "non-empty supported model identifier", "may require model download and available memory"),
    "device_id": ("integer or string", "device identifier", "Target accelerator or device identifier.", "workload-specific valid device", "must refer to a visible device"),
    "output_format": ("string", "format token", "Format used for benchmark output.", "one of the workload-supported formats", "must agree with parser expectations"),
    "prompt_source": ("string", "source identifier", "Source or generation mode for input prompts.", "supported source token", "determines whether prompts are synthetic or dataset-backed"),
    "seed": ("integer", "seed value", "Random seed used for reproducibility.", "integer >= 0", "same seed does not guarantee identical hardware timing"),
    "input_len": ("integer", "tokens", "Input or prompt token length.", "integer >= 0; workload-specific maximum", "distinct from total sequence length"),
    "num_gpus": ("integer", "GPUs", "Number of GPUs participating in the workload.", "integer >= 1 and <= visible GPUs", "may constrain tensor parallelism and process-grid dimensions"),
    "num_threads": ("integer", "threads", "Execution threads issuing work.", "integer >= 1", "distinct from num_workers"),
    "output_len": ("integer", "tokens", "Number of generated or output tokens.", "integer >= 0", "input_len + output_len may be limited by model context"),
    "precision": ("string", "format token", "Arithmetic or compute precision.", "workload-specific supported precision token", "may differ from storage dtype"),
    "dataset_name": ("string", "dataset identifier", "Dataset or prompt corpus identifier.", "non-empty supported dataset identifier", "record canonical dataset and split when applicable"),
    "max_concurrency": ("integer", "in-flight requests", "Maximum client-side requests in flight.", "integer >= 1", "client queue limit; distinct from server and engine capacity"),
    "page_size": ("integer or string", "bytes or page token", "Memory page size used by the workload.", "supported OS or benchmark page size", "units must be explicit because values may be written as 4K, 2M, or bytes"),
    "request_rate": ("number", "requests per second", "Rate at which requests or prompts are submitted.", "number >= 0", "may be ignored or overridden when closed-loop concurrency is used"),
    "tensor_parallel_size": ("integer", "GPUs", "Number of GPUs across which model tensors are sharded.", "integer >= 1 and <= visible GPUs", "must be compatible with num_gpus and backend support"),
    "thread_affinity": ("string", "affinity specification", "Policy or CPU/NUMA placement for execution threads.", "supported affinity expression", "interacts with NUMA placement and reproducibility"),
    "block_size": ("integer or string", "workload-specific", "Work unit or block size.", "positive workload-specific value", "meaning varies: I/O bytes, GPU threads, matrix tile, or panel size"),
    "duration": ("number", "seconds", "Wall-clock or sampling duration.", "number > 0", "must define precedence or mutual exclusion with num_iterations"),
    "ecc_check": ("boolean", "boolean", "Whether ECC error telemetry is collected or validated.", "true or false", "may require supported hardware telemetry"),
    "gpu_memory_utilization": ("number", "fraction", "Fraction of GPU memory the runtime may use.", "0 < value <= 1", "must leave sufficient memory for runtime overhead"),
    "num_prompts": ("integer", "prompts", "Total prompts or requests submitted.", "integer >= 1", "interacts with request_rate and concurrency"),
    "sequence_len": ("integer", "tokens", "Total input sequence or context length.", "integer >= 1; workload-specific maximum", "do not confuse with prompt-only input_len"),
    "access_pattern": ("string", "pattern token", "Memory traversal or access pattern.", "supported pattern token", "affects locality and measured bandwidth/latency"),
    "array_size": ("integer", "elements", "Number of elements in benchmark arrays.", "integer > 0 and memory-feasible", "must fit target memory and may need block alignment"),
    "gradient_accumulation": ("integer", "forward passes", "Number of forward/backward passes accumulated per optimizer update.", "integer >= 1", "changes effective batch size and optimizer-step frequency"),
    "image_size": ("integer or tuple", "pixels", "Spatial image dimensions.", "positive supported dimensions", "must match model and memory limits"),
    "K": ("integer", "matrix dimension", "GEMM K matrix dimension.", "integer > 0", "M, N, and K jointly determine memory and compute cost"),
    "learning_rate": ("number", "optimizer units", "Optimizer step size.", "number > 0; workload-specific range", "interacts with optimizer and gradient_accumulation"),
    "max_bytes": ("integer", "bytes", "Maximum message or transfer size.", "integer >= min_bytes", "must be >= min_bytes and memory-feasible"),
    "max_num_seqs": ("integer", "sequences", "Maximum sequences the inference engine processes concurrently.", "integer >= 1", "engine capacity; distinct from client concurrency"),
    "max_running_requests": ("integer", "requests", "Maximum server-side requests actively running.", "integer >= 1", "server capacity; excess requests queue or are rejected"),
    "max_total_tokens": ("integer", "tokens", "Maximum server-side tokens held in the scheduling or KV-cache pool.", "integer >= 1", "must accommodate configured input/output lengths"),
    "min_bytes": ("integer", "bytes", "Minimum message or transfer size.", "integer > 0", "must be <= max_bytes"),
    "M": ("integer", "matrix dimension", "GEMM M matrix dimension.", "integer > 0", "M, N, and K jointly determine memory and compute cost"),
    "N": ("integer", "matrix dimension", "GEMM N matrix dimension.", "integer > 0", "M, N, and K jointly determine memory and compute cost"),
    "num_workers": ("integer", "workers", "Worker processes or data-loader/service workers.", "integer >= 0", "distinct from execution threads"),
    "numa_node": ("integer", "NUMA node", "NUMA node selected for placement.", "integer >= 0 and present on host", "interacts with thread_affinity and memory placement"),
    "optimizer": ("string", "optimizer token", "Optimization algorithm used for training.", "supported optimizer token", "interacts with learning_rate and gradient_accumulation"),
    "schedule_policy": ("string", "policy token", "Request scheduling policy.", "supported backend policy token", "affects queueing, fairness, and latency"),
    "step_factor": ("number", "multiplicative factor", "Multiplier between successive sweep sizes.", "number > 1", "controls logarithmic sweep progression"),
    "stride": ("integer", "workload-specific units", "Distance between successive memory accesses.", "positive workload-specific value", "units must be documented as bytes, elements, or tokens"),
    "working_set_size": ("integer or string", "bytes or size token", "Memory region size used by the workload.", "positive memory-feasible value", "must fit target memory and may interact with page_size and stride"),
}

LABEL_PATTERN = re.compile(
    r"(?<![A-Za-z0-9_])(?P<name>[A-Za-z_][A-Za-z0-9_]*)\s*:\s*"
    r"(?P<meaning>.*?)(?=\.\s+[A-Za-z_][A-Za-z0-9_]*\s*:|$)",
    re.DOTALL,
)


def extract_meaning(description: str, parameter: str) -> str:
    pattern = re.compile(
        rf"(?<![A-Za-z0-9_]){re.escape(parameter)}\s*:\s*"
        r"(?P<meaning>.*?)(?=\.\s+[A-Za-z_][A-Za-z0-9_]*\s*:|$)",
        re.DOTALL,
    )
    match = pattern.search(description)
    if not match:
        return "Defined by the workload description."
    return match.group("meaning").strip().rstrip(".")


def validate_active_parameters(worksheet) -> None:
    """Ensure every active parameter is named in its workload description."""
    missing: list[str] = []
    parameter_columns = _parameter_columns(worksheet)
    description_column = _description_column(worksheet)
    for row in range(2, worksheet.max_row + 1):
        workload_id = worksheet.cell(row, 1).value
        if workload_id is None:
            continue
        try:
            int(workload_id)
        except (TypeError, ValueError):
            # Repeated Workload_Definitions header rows are not workloads.
            continue
        description = str(worksheet.cell(row, description_column).value or "")
        for column in parameter_columns:
            parameter = worksheet.cell(row, column).value
            if parameter in (None, "—"):
                continue
            name = str(parameter).strip()
            if not re.search(
                rf"(?<![A-Za-z0-9_]){re.escape(name)}(?![A-Za-z0-9_])",
                description,
            ):
                missing.append(f"{workload_id}:{name}")
    if missing:
        raise SystemExit(
            "[FAIL] Active workbook parameters missing from their execution descriptions: "
            + ", ".join(missing)
        )


def build_catalog() -> dict:
    workbook = load_workbook(WORKBOOK, read_only=True, data_only=False)
    worksheet = _active_sheet(workbook)
    validate_active_parameters(worksheet)
    workload_names: dict[int, str] = {}
    uses: defaultdict[str, list[int]] = defaultdict(list)
    meanings: defaultdict[str, dict[int, str]] = defaultdict(dict)
    parameter_columns = _parameter_columns(worksheet)
    description_column = _description_column(worksheet)

    for row in range(2, worksheet.max_row + 1):
        workload_id = worksheet.cell(row, 1).value
        if workload_id is None:
            continue
        try:
            workload_id = int(workload_id)
        except (TypeError, ValueError):
            continue
        workload_name = worksheet.cell(row, 2).value
        description = worksheet.cell(row, description_column).value or ""
        workload_names[workload_id] = workload_name
        active_parameters = {
            worksheet.cell(row, column).value
            for column in parameter_columns
            if worksheet.cell(row, column).value not in (None, "—")
        }
        for parameter in active_parameters:
            uses[parameter].append(workload_id)
            meanings[parameter][workload_id] = extract_meaning(description, parameter)

    entries = {}
    for parameter in sorted(uses, key=lambda item: (-len(uses[item]), item)):
        if len(uses[parameter]) < 2:
            continue
        data_type, units, description, valid_range, interaction = CATALOG_METADATA.get(
            parameter,
            (
                "string",
                "workload-specific",
                "Shared workload parameter.",
                "workload-specific",
                "workload-specific interaction rules",
            ),
        )
        workload_ids = sorted(uses[parameter])
        entries[parameter] = {
            "description": description,
            "data_type": data_type,
            "units": units,
            "default_value": "workload-specific",
            "valid_range": valid_range,
            "required": "workload-specific",
            "semantic_scope": "workload-specific; see workload definitions below",
            "workloads_using": workload_ids,
            "workload_specific_overrides": {
                str(workload_id): {
                    "workload_name": workload_names[workload_id],
                    "meaning": meanings[parameter][workload_id],
                }
                for workload_id in workload_ids
            },
            "cli_config_key": "workload-specific",
            "interaction_precedence": interaction,
            "notes_constraints": (
                "Do not assume this shared name has identical units or semantics across all workloads."
            ),
        }
    return {
        "schema_version": "1.0.0",
        "purpose": "Canonical definitions for parameters shared by at least two workloads.",
        "source_workbook": f"BenchmarkSpecDefinitions.xlsx::{worksheet.title}",
        "source_of_truth": (
            "Workload-specific parameter placement and execution descriptions remain authoritative "
            "in the workbook; this catalog standardizes shared terminology and exposes semantic overrides."
        ),
        "generation_note": (
            "workloads_using and workload_specific_overrides are derived from the workbook and should "
            "be regenerated when parameter cells or descriptions change."
        ),
        "parameters": entries,
    }


def main() -> None:
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--check",
        action="store_true",
        help="Validate workbook parameter references without writing the catalog.",
    )
    args = parser.parse_args()
    catalog = build_catalog()
    if args.check:
        print(f"[PASS] Validated {len(catalog['parameters'])} shared parameters.")
        return
    with OUTPUT.open("w", encoding="utf-8", newline="\n") as stream:
        yaml.safe_dump(catalog, stream, sort_keys=False, allow_unicode=True, width=120)
    print(f"Wrote {OUTPUT} with {len(catalog['parameters'])} shared parameters.")


if __name__ == "__main__":
    main()
